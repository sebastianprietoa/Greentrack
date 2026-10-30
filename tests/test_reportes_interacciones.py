import io
from openpyxl import load_workbook
import pytest

from tests.test_integridad_postgres import pg
from tests.test_combustibles_carbono import fuel_pg
from tests.test_agua_coherencia import water_pg,vista


@pytest.mark.parametrize('periodo',['Todos',''])
def test_ghg_historico_no_descarta_registros(water_pg,periodo):
    cursor=water_pg.cursor()
    for year,value in [(2025,100),(2026,200)]:
        cursor.execute("INSERT INTO registros (empresa,fecha,fuente,alcance,categoria,emision) VALUES ('AUDIT_A',%s,'Combustión Fija','Alcance 1','Diésel',%s)",(f'{year}-01-01',value))
    water_pg.commit()
    response,_=vista(water_pg,'/exportar_ghg?anio='+periodo)
    assert response.status_code==200
    book=load_workbook(io.BytesIO(response.data),read_only=True,data_only=True)
    assert list(book['Resumen GHG'].values)[-1][1]==300
    assert book['Alcance 1'].max_row==3
    book.close()


def test_dashboard_sin_medida_no_muestra_intensidad_cero(water_pg):
    response,context=vista(water_pg,'/dashboard?anio=Todos')
    assert response.status_code==200
    assert context['intensidad_emisiones'] is None
    assert '0.000000' not in response.get_data(as_text=True)
    assert 'exportar_ghg?anio=Todos' in response.get_data(as_text=True)


@pytest.mark.parametrize('path,destino',[
    ('/agua/registro','/registro?tab=agua'),('/agua/reporte','/agua/resultados')])
def test_accesos_historicos_conducen_a_funciones_activas(water_pg,path,destino):
    response,_=vista(water_pg,path)
    assert response.status_code==302
    assert response.location==destino


def test_ghg_rechaza_anio_invalido(water_pg):
    response,_=vista(water_pg,'/exportar_ghg?anio=incorrecto')
    assert response.status_code==302
    assert response.location=='/exportar'


def test_exportar_desde_agua_conserva_modulo_y_periodo(water_pg):
    from html.parser import HTMLParser
    class Enlaces(HTMLParser):
        def __init__(self):
            super().__init__()
            self.urls=[]
        def handle_starttag(self,tag,attrs):
            attributes=dict(attrs)
            if tag=='a' and attributes.get('class')=='wf-action':
                self.urls.append(attributes.get('href'))
    response,_=vista(water_pg,'/agua?periodo=2026-01')
    assert response.status_code==200
    enlaces=Enlaces()
    enlaces.feed(response.get_data(as_text=True))
    assert enlaces.urls==['/agua/resultados/descargar?periodo=2026-01']
    download,_=vista(water_pg,enlaces.urls[0])
    assert download.status_code==200
    book=load_workbook(io.BytesIO(download.data),read_only=True)
    assert 'Flujos de agua' in book.sheetnames
    assert 'Resumen' in book.sheetnames
    book.close()
