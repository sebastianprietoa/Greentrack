"""Períodos consistentes en paneles, navegación y reportes; tablas temporales."""
import io
from html.parser import HTMLParser

from openpyxl import load_workbook
import pytest

from tests.test_aceptacion_roles import acceptance_pg, water_pg, fuel_pg, pg, database, client_for, context_get


@pytest.fixture
def periodos(acceptance_pg):
    conn, _ = acceptance_pg
    cursor = conn.cursor()
    cursor.execute("UPDATE usuarios SET anio_default='2025' WHERE id=3")
    for company, year, value in [('AUDIT_A',2025,10),('AUDIT_A',2026,100),('AUDIT_B',2025,999)]:
        for source, scope in [('Combustión Fija','Alcance 1'),('Electricidad','Alcance 2'),('Residuos','Alcance 3'),('Refrigerantes','Alcance 1')]:
            cursor.execute('''INSERT INTO registros (empresa,fecha,fuente,alcance,categoria,actividad,cantidad,unidad,emision,emision_ubicacion,sistema,tiene_irec)
                VALUES (%s,%s,%s,%s,'Prueba','Prueba',%s,'kg',%s,%s,'SEN','No')''',
                (company,f'{year}-01-01',source,scope,value,value,value/2))
    conn.commit()
    return conn


class Enlaces(HTMLParser):
    def __init__(self, html):
        super().__init__()
        self.urls=[]
        self.feed(html)
    def handle_starttag(self, tag, attrs):
        if tag=='a':
            self.urls.append(dict(attrs).get('href',''))


@pytest.mark.parametrize('path,key,series',[
    ('/combustion','total_emision','meses_comb'),
    ('/electricidad','total_mercado','tendencia_elec'),
    ('/residuos','total_emision','datos_mes'),
    ('/refrigerantes','total_refrigerantes','datos_meses'),
])
@pytest.mark.parametrize('query,expected,years',[
    ('',10,{'2025'}),('?anio=2026',100,{'2026'}),
    ('?anio=Todos',110,{'2025','2026'}),('?anio=2024',0,set()),
    ('?anio=incorrecto',110,{'2025','2026'}),
])
def test_panel_filtra_total_series_y_tablas(periodos,path,key,series,query,expected,years):
    with database(periodos):
        response,context=context_get(client_for('empresa'),path+query)
    assert float(context[key])==expected
    rows=context[series]
    assert {str(r if isinstance(r,str) else r['mes'])[:4] for r in rows}==years
    for row in context.get('registros',[]):
        assert str(row['fecha'])[:4] in years
    if path=='/electricidad':
        assert context['total_kwh']==expected
        assert context['total_ubicacion']==expected/2
    if path=='/residuos':
        assert sum(r['total'] for r in context['datos_categoria'])==expected
        assert sum(r['total'] for r in context['datos_tratamiento'])==expected
    assert 'id="anio-panel"' in response.get_data(as_text=True)


def test_historial_y_exportacion_usan_mismo_anio(periodos):
    with database(periodos):
        client=client_for('empresa')
        for query,year,total in [('', '2025',4),('?anio=2026','2026',4),('?anio=','Todos',8),('?anio=Todos','Todos',8)]:
            response,context=context_get(client,'/historial'+query)
            assert context['total']==total
            url=f'/exportar_ghg?anio={year}'
            assert url in Enlaces(response.get_data(as_text=True)).urls
            report=client.get(url)
            book=load_workbook(io.BytesIO(report.data),read_only=True,data_only=True)
            assert list(book['Resumen GHG'].values)[-1][1]==(40 if year=='2025' else 400 if year=='2026' else 440)
            book.close()


def test_exportar_sin_parametro_ghg_respeta_preferencia(periodos):
    with database(periodos):
        client=client_for('empresa')
        response,context=context_get(client,'/exportar')
        assert context['anio_seleccionado']=='2025'
        assert '/exportar_ghg?anio=2025' in Enlaces(response.get_data(as_text=True)).urls
        report=client.get('/exportar_ghg')
    book=load_workbook(io.BytesIO(report.data),read_only=True,data_only=True)
    assert list(book['Resumen GHG'].values)[-1][1]==40
    book.close()


@pytest.mark.parametrize('periodo',['2026-01','2026',''])
def test_agua_conserva_periodo_en_registros_exportacion_y_regreso(periodos,periodo):
    with database(periodos):
        client=client_for('empresa')
        response,context=context_get(client,'/agua?periodo='+periodo)
        html=response.get_data(as_text=True)
        assert context['periodo']==(periodo or None)
        if periodo:
            assert f'value="{periodo}" selected' in html
        urls=Enlaces(html).urls
        assert '/combustion?anio='+(periodo[:4] if periodo else 'Todos') in urls
        assert '/agua/resultados?periodo='+periodo in urls
        response,context=context_get(client,'/agua/resultados?periodo='+periodo)
        urls=Enlaces(response.get_data(as_text=True)).urls
        assert '/agua?periodo='+periodo in urls
        assert '/agua/resultados/descargar?periodo='+periodo+'&vista=mensual' in urls


def test_navegacion_conserva_anio_y_agua_ofrece_anual(periodos):
    with database(periodos):
        client=client_for('empresa')
        response,_=context_get(client,'/electricidad?anio=2026')
        urls=Enlaces(response.get_data(as_text=True)).urls
        for path in ['dashboard','combustion','electricidad','refrigerantes','residuos','exportar']:
            assert f'/{path}?anio=2026' in urls
        assert '/agua?periodo=2026' in urls
        response,context=context_get(client,'/agua')
        assert context['periodo']=='2025'
        assert '2026' in context['periodos_disponibles']
        assert 'value="2025" selected' in response.get_data(as_text=True)
