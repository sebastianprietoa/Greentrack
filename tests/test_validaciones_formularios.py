"""Validaciones de formularios y coherencia de mensajes sin escritura productiva."""
import io
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch

from openpyxl import Workbook,load_workbook
import pytest

from services.agua_inventario import MENSAJES,crear_resumen_agua,validar_nuevo_flujo
from tests.test_aceptacion_roles import acceptance_pg,water_pg,fuel_pg,pg,database,client_for,context_get


def mensajes(client):
    with client.session_transaction() as state:
        return ' '.join(text for _,text in state.get('_flashes',[]))


@pytest.mark.parametrize('extra,fragment',[
    ({'fecha_inicio':'2025-12-31','fecha_fin':'2025-01-01'},'posterior'),
    ({'fecha_inicio':'2025-02-30'},'fechas válidas'),
    ({'fecha_fin':'no-es-fecha'},'fechas válidas'),
    ({'fecha_inicio':'2025-1-01'},'fechas válidas'),
    ({'anio':'20255'},'año válido'),
    ({'anio':'0000'},'año válido'),
])
def test_exportacion_rechaza_filtros_invalidos_antes_de_consultar(extra,fragment):
    import app as module
    client=client_for('empresa')
    args=dict(anio='2025',fecha_inicio='',fecha_fin='',alcance=['Alcance 1','Alcance 2'],fuente=['Electricidad','Residuos'])
    args.update(extra)
    with patch.object(module,'get_db',side_effect=AssertionError('No debe consultar filtros inválidos')):
        response=client.get('/exportar_avanzado',query_string=args)
    assert response.status_code==302
    kept=parse_qs(urlparse(response.location).query,keep_blank_values=True)
    for key,value in args.items():
        assert kept[key]==(value if isinstance(value,list) else [value])
    assert fragment in mensajes(client)


@pytest.mark.parametrize('start,end,year,expected',[
    ('2025-01-01','2025-01-01','2025',1),
    ('2025-01-01','2025-02-28','2025',2),
    ('2025-01-01','','2025',2),
    ('','2025-01-31','2025',1),
    ('','','Todos',3),
])
def test_exportacion_combina_filtros_y_admite_limites_inclusivos(acceptance_pg,start,end,year,expected):
    conn,_=acceptance_pg
    cursor=conn.cursor()
    for date in ['2025-01-01','2025-02-01','2026-01-01']:
        cursor.execute("INSERT INTO registros (empresa,fecha,fuente,cantidad,unidad,factor,emision) VALUES ('AUDIT_A',%s,'Electricidad',100,'kWh',0.3,30)",(date,))
    conn.commit()
    with database(conn):
        response=client_for('empresa').get('/exportar_avanzado',query_string=dict(anio=year,fecha_inicio=start,fecha_fin=end))
    assert response.status_code==200
    book=load_workbook(io.BytesIO(response.data),read_only=True)
    assert book['Registros'].max_row==expected+1
    book.close()


def test_sin_resultados_conserva_fechas_y_selecciones(acceptance_pg):
    conn,_=acceptance_pg
    cursor=conn.cursor()
    cursor.execute("INSERT INTO registros (empresa,fecha,fuente,alcance,emision) VALUES ('AUDIT_A','2025-01-01','Electricidad','Alcance 2',30)")
    conn.commit()
    args=dict(anio='2026',fecha_inicio='2025-01-01',fecha_fin='2025-12-31',alcance='Alcance 2',fuente='Electricidad')
    with database(conn):
        client=client_for('empresa')
        response=client.get('/exportar_avanzado',query_string=args)
        assert response.status_code==302
        assert parse_qs(urlparse(response.location).query)=={k:[v] for k,v in args.items()}
        response,context=context_get(client,response.location)
    html=response.get_data(as_text=True)
    assert 'value="2025-01-01"' in html and 'value="2025-12-31"' in html
    assert context['anio_seleccionado']=='2026'
    assert 'value="Alcance 2" checked' in html
    assert 'value="Electricidad" checked' in html
    assert 'se combinan' in html
    assert html.count('No hay registros que coincidan con los filtros seleccionados.')==1


def excel(sheet='Datos Centralizados',headers=('Cantidad',),rows=()):
    book=Workbook()
    book.active.title=sheet
    book.active.append(list(headers))
    for row in rows: book.active.append(list(row))
    output=io.BytesIO();book.save(output);output.seek(0)
    return output


@pytest.mark.parametrize('payload,fragment',[
    (None,'No se seleccionó'),
    ((io.BytesIO(b''),'datos.csv'),'archivo .xlsx'),
    ((io.BytesIO(b''),'datos.xls'),'archivo .xlsx'),
    ((io.BytesIO(b''),'datos.xlsx'),'No se pudo leer'),
    ((io.BytesIO(b'esto no es Excel'),'datos.xlsx'),'No se pudo leer'),
])
def test_importacion_archivo_invalido_no_genera_importacion(acceptance_pg,payload,fragment):
    conn,_=acceptance_pg
    with database(conn):
        client=client_for('empresa')
        response=client.post('/importar',data={'archivo':payload} if payload else {},content_type='multipart/form-data')
    assert response.status_code==302
    assert fragment in mensajes(client)
    with client.session_transaction() as state:
        assert 'import_preview' not in state
    cursor=conn.cursor();cursor.execute('SELECT COUNT(*) FROM registros');assert cursor.fetchone()[0]==0


@pytest.mark.parametrize('kind,fragment',[
    ('pestana','Datos Centralizados'),('columna','Cantidad'),('vacio','no contiene registros'),
])
def test_importacion_estructura_invalida_da_mensaje_claro(acceptance_pg,kind,fragment):
    conn,_=acceptance_pg
    file=excel(sheet='Otra') if kind=='pestana' else excel(headers=('Otra',),rows=((1,),)) if kind=='columna' else excel()
    with database(conn):
        client=client_for('empresa')
        response=client.post('/importar',data={'archivo':(file,'datos.XLSX')})
    assert response.status_code==302
    assert fragment in mensajes(client)
    assert 'Error técnico' not in mensajes(client)


@pytest.mark.parametrize('entrada,salida,estado',[('no','no','sin_medicion'),('boleta','no','solo_entrada')])
def test_validaciones_de_agua_y_resumen_comparten_mensaje_azul(acceptance_pg,entrada,salida,estado):
    conn,_=acceptance_pg
    cursor=conn.cursor()
    cursor.execute('INSERT INTO configuracion (empresa,medicion_entrada_agua,medicion_salida_agua) VALUES (%s,%s,%s)',('AUDIT_A',entrada,salida))
    conn.commit()
    with pytest.raises(ValueError,match='huella azul') as error:
        validar_nuevo_flujo(cursor,'AUDIT_A',101,'2025-01','retorno',1)
    assert str(error.value)==MENSAJES[estado]
    assert crear_resumen_agua(conn,'AUDIT_A')['mensaje_medicion_agua']==MENSAJES[estado]


def test_formulario_excel_ofrece_error_accesible(acceptance_pg):
    conn,_=acceptance_pg
    with database(conn):response,_=context_get(client_for('empresa'),'/importar')
    html=response.get_data(as_text=True)
    assert 'id="archivoError" role="alert"' in html
    assert 'aria-label="Archivo Excel"' in html
    assert 'accept=".xlsx"' in html
