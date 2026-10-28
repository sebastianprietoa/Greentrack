"""Aceptación por roles y recorridos completos en tablas temporales de staging."""
import io
import secrets
from contextlib import contextmanager
from unittest.mock import patch

from flask import template_rendered
from openpyxl import load_workbook
import pytest

from auth_utils import hash_password
from services.migracion_integridad import reconciliar_integridad
from tests.test_integridad_postgres import pg,SharedConnection
from tests.test_combustibles_carbono import fuel_pg
from tests.test_agua_coherencia import water_pg


@pytest.fixture
def acceptance_pg(water_pg):
    cursor=water_pg.cursor()
    tables=['usuarios','pdf_uploads','pending_pdf_uploads','tickets','agua_consumo','agua_cuencas',
        'agua_afluentes','agua_costos','residuos_registros','irec_certificados','energeticos_empresa']
    for table in tables:
        cursor.execute(f'CREATE TEMP TABLE {table} (LIKE public.{table} INCLUDING DEFAULTS INCLUDING INDEXES)')
        cursor.execute(f'CREATE TEMP SEQUENCE audit_{table}_seq START 5000')
        cursor.execute(f"ALTER TABLE {table} ALTER COLUMN id SET DEFAULT nextval('pg_temp.audit_{table}_seq')")
    cursor.execute('DROP TABLE pg_temp.vehiculos')
    cursor.execute('CREATE TEMP TABLE vehiculos (LIKE public.vehiculos INCLUDING DEFAULTS INCLUDING INDEXES)')
    cursor.execute('CREATE TEMP SEQUENCE audit_vehiculos_seq START 5000')
    cursor.execute("ALTER TABLE vehiculos ALTER COLUMN id SET DEFAULT nextval('pg_temp.audit_vehiculos_seq')")
    cursor.execute("INSERT INTO vehiculos (id,empresa,patente,tipo) VALUES (401,'AUDIT_A','AUDIT01','auto'),(402,'AUDIT_B','AUDIT02','auto')")
    cursor.execute('DROP TABLE pg_temp.factores')
    cursor.execute('CREATE TEMP TABLE factores (LIKE public.factores INCLUDING DEFAULTS INCLUDING INDEXES)')
    cursor.execute("INSERT INTO factores (categoria,unidad,factor,anio) VALUES ('Diésel fósil puro','Litros',2.66155,0)")
    password=secrets.token_urlsafe(18)
    hashed=hash_password(password)
    for id,company,role in [(1,'AUDIT_ADMIN','admin'),(2,'AUDIT_MANAGER','gestor_empresa'),(3,'AUDIT_A','empresa'),(4,'AUDIT_B','empresa'),(5,'AUDIT_OTHER_MANAGER','gestor_empresa')]:
        cursor.execute('INSERT INTO usuarios (id,empresa,email,password,es_admin,rol,es_demo) VALUES (%s,%s,%s,%s,%s,%s,false)',
            (id,company,f'{company.lower()}@example.invalid',hashed,int(role=='admin'),role))
    water_pg.commit()
    return water_pg,password


def client_for(role=None):
    import app as module
    module.app.config['TESTING']=True
    client=module.app.test_client()
    if role:
        id,empresa={'admin':(1,'AUDIT_ADMIN'),'gestor_empresa':(2,'AUDIT_MANAGER'),'empresa':(3,'AUDIT_A')}[role]
        with client.session_transaction() as session:
            session.update(user_id=id,empresa=empresa,es_admin=int(role=='admin'),rol=role)
    return client


@contextmanager
def database(connection):
    import app as module
    import routes.auth_routes as auth
    with patch.object(module,'get_db',return_value=SharedConnection(connection)),patch.object(auth,'get_db',return_value=SharedConnection(connection)):
        yield


def context_get(client,path):
    import app as module
    contexts=[]
    def capture(sender,template,context,**extra): contexts.append(context)
    with template_rendered.connected_to(capture,module.app):
        response=client.get(path)
    assert response.status_code==200
    return response,contexts[-1]


@pytest.mark.parametrize('role',[None,'empresa','gestor_empresa'])
def test_admin_only_operations_reject_other_roles_before_database(role):
    import app as module
    client=client_for(role)
    cases=[('GET','/admin/dashboard'),('GET','/admin/factores'),('POST','/admin/cargar_factores'),
        ('POST','/admin/cargar_electricidad'),('POST','/admin/preview_electricidad'),
        ('POST','/admin/eliminar_empresa/3'),('GET','/admin/exportar_todo'),('GET','/admin/exportar_ghg_global'),
        ('GET','/admin/huella-hidrica'),('POST','/admin/huella-hidrica/sedes'),
        ('POST','/admin/huella-hidrica/sedes/101/toggle'),('POST','/admin/huella-hidrica/factores'),
        ('POST','/admin/huella-hidrica/factores/1/toggle'),('POST','/admin/huella-hidrica/factores/preview'),
        ('POST','/admin/huella-hidrica/factores/cargar'),('GET','/admin/pendientes'),
        ('POST','/admin/aprobar_pdf/1'),('POST','/admin/rechazar_pdf/1'),('GET','/admin/tickets'),
        ('POST','/admin/tickets/1/responder')]
    with patch.object(module,'get_db',side_effect=AssertionError('Una ruta denegada no debe acceder a datos')):
        for method,path in cases:
            response=client.open(path,method=method)
            assert response.status_code==302,path


@pytest.mark.parametrize('role',['admin','gestor_empresa','empresa'])
def test_login_web_api_and_global_api_permissions(acceptance_pg,role):
    connection,password=acceptance_pg
    import app as module
    names={'admin':'audit_admin','gestor_empresa':'audit_manager','empresa':'audit_a'}
    destinations={'admin':'/admin/dashboard','gestor_empresa':'/admin/empresas','empresa':'/dashboard'}
    with database(connection):
        client=client_for()
        response=client.post('/',data={'email':names[role]+'@example.invalid','password':password})
        assert response.status_code==302
        assert response.location==destinations[role]
        api=client.post('/api/login',json={'email':names[role]+'@example.invalid','password':password})
        assert api.status_code==200
        token=api.json['access_token']
        response=client.get('/api/admin/emisiones-por-empresa',headers={'Authorization':'Bearer '+token})
        assert response.status_code==(200 if role=='admin' else 403)
        assert client.get('/api/admin/emisiones-por-empresa',headers={'Authorization':'Bearer invalid'}).status_code==401
        client.get('/logout')
        with client.session_transaction() as session: assert 'user_id' not in session


def test_legacy_admin_role_cannot_remove_admin_permissions(acceptance_pg):
    connection,password=acceptance_pg
    cursor=connection.cursor()
    cursor.execute("UPDATE usuarios SET rol='empresa' WHERE email='audit_admin@example.invalid'")
    connection.commit()
    with database(connection):
        client=client_for()
        response=client.post('/',data={'email':'audit_admin@example.invalid','password':password})
        assert response.location=='/admin/dashboard'
        with client.session_transaction() as session:
            assert session['es_admin']==1
            assert session['rol']=='admin'
        assert client.get('/admin/empresas').status_code==200
        assert client.post('/admin/crear_empresa',data={}).location=='/admin/empresas'
        api=client.post('/api/login',json={'email':'audit_admin@example.invalid','password':password})
        assert api.status_code==200
        assert api.json['user']['rol']=='admin'


def test_admin_with_zero_id_can_manage_companies(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        client=client_for()
        with client.session_transaction() as session:
            session.update(user_id=0,empresa='ADMIN',es_admin=1,rol='admin')
        assert client.get('/admin/empresas').status_code==200
        assert client.post('/admin/crear_empresa',data={}).location=='/admin/empresas'


@pytest.mark.parametrize('role',['admin','gestor_empresa'])
def test_allowed_company_management_and_administration_pages(acceptance_pg,role):
    connection,_=acceptance_pg
    seed_consumption_and_support(connection)
    paths=['/admin/empresas','/admin/empresa/AUDIT_A','/admin/exportar/AUDIT_A']
    if role=='admin':
        paths+=['/admin/dashboard','/admin/huella-hidrica','/admin/factores','/admin/pendientes','/admin/tickets']
    with database(connection):
        client=client_for(role)
        for path in paths:
            assert client.get(path).status_code==200,path
        _,listing=context_get(client,'/admin/empresas')
        assert listing['total_empresas']==len(listing['empresas'])==2


@pytest.mark.parametrize('role',['admin','gestor_empresa'])
def test_create_company_form_rejects_empty_input(acceptance_pg,role):
    connection,_=acceptance_pg
    with database(connection):
        assert client_for(role).post('/admin/crear_empresa',data={}).status_code==302
    cursor=connection.cursor()
    cursor.execute('SELECT COUNT(*) FROM usuarios')
    assert cursor.fetchone()[0]==5


def seed_consumption_and_support(connection):
    cursor=connection.cursor()
    for company,id in [('AUDIT_A',101),('AUDIT_B',102)]:
        cursor.execute("INSERT INTO registros (id,empresa,fecha,fuente,alcance,categoria,unidad,cantidad,factor,emision) VALUES (%s,%s,'2026-01-01','Combustible Móvil','Alcance 1','Diésel','litros',100,2.68,268)",(id,company))
        cursor.execute('INSERT INTO combustible_movil (empresa,vehiculo_id,registro_id,periodo) VALUES (%s,%s,%s,%s)',(company,401 if id==101 else 402,id,'2026-01'))
        cursor.execute('INSERT INTO resultados_huella_suministro_agua (empresa,registro_id,huella_suministro_m3) VALUES (%s,%s,2)',(company,id))
        cursor.execute("INSERT INTO tickets (empresa,asunto,descripcion,fecha_creacion) VALUES (%s,'Prueba','Soporte','2026-01-01')",(company,))
        cursor.execute("INSERT INTO pending_pdf_uploads (empresa,nombre_archivo,estado) VALUES (%s,'test.pdf','pendiente')",(company,))
    connection.commit()


@pytest.mark.parametrize('role',['admin','gestor_empresa'])
def test_company_rename_preserves_linked_consumption_and_support(acceptance_pg,role):
    connection,password=acceptance_pg
    seed_consumption_and_support(connection)
    reconciliar_integridad(connection)
    connection.commit()
    with database(connection):
        response=client_for(role).post('/admin/editar_empresa/3',data={'empresa':'AUDIT_RENAMED','email':'audit_a@example.invalid'})
        assert response.status_code==302
    cursor=connection.cursor()
    for table in ['usuarios','registros','combustible_movil','resultados_huella_suministro_agua','tickets','pending_pdf_uploads']:
        cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE empresa='AUDIT_RENAMED'")
        assert cursor.fetchone()[0]==1,table
        cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE empresa='AUDIT_B'")
        assert cursor.fetchone()[0]==1,table
    with database(connection):
        client=client_for()
        assert client.post('/',data={'email':'audit_a@example.invalid','password':password}).location=='/dashboard'
        _,dashboard=context_get(client,'/dashboard?anio=2026')
        assert dashboard['total_emision']==268


def test_delete_company_removes_pending_support_only_for_target(acceptance_pg):
    connection,_=acceptance_pg
    seed_consumption_and_support(connection)
    reconciliar_integridad(connection)
    connection.commit()
    with database(connection):
        assert client_for('admin').post('/admin/eliminar_empresa/3').status_code==302
    cursor=connection.cursor()
    for table in ['usuarios','registros','combustible_movil','resultados_huella_suministro_agua','tickets','pending_pdf_uploads']:
        cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE empresa='AUDIT_A'")
        assert cursor.fetchone()[0]==0,table
        cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE empresa='AUDIT_B'")
        assert cursor.fetchone()[0]==1,table


@pytest.mark.parametrize('operation',['editar_empresa','resetear_password'])
def test_manager_cannot_modify_another_manager(acceptance_pg,operation):
    connection,password=acceptance_pg
    with database(connection):
        response=client_for('gestor_empresa').post('/admin/'+operation+'/5',data={
            'empresa':'CHANGED','email':'changed@example.invalid','nueva_password':secrets.token_urlsafe(20)})
        assert response.status_code==302
    cursor=connection.cursor()
    cursor.execute('SELECT empresa,password FROM usuarios WHERE id=5')
    company,hashed=cursor.fetchone()
    from auth_utils import verify_password
    assert company=='AUDIT_OTHER_MANAGER'
    assert verify_password(password,hashed)


def test_anonymous_ticket_urls_lead_to_existing_login():
    client=client_for()
    for path in ['/tickets','/admin/tickets']:
        response=client.get(path,follow_redirects=True)
        assert response.status_code==200
        assert b'Iniciar' in response.data


def test_company_consumption_dashboard_filter_export_and_delete(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        company=client_for('empresa')
        response=company.post('/api/combustible/movil',json={'registros':[{
            'vehiculo_id':401,'periodo':'2026-01','cantidad':100,'unidad':'litros','combustible':'diesel'}]})
        assert response.status_code==200
        assert response.json['guardados']==1
        assert company.post('/configuracion',data={'medida_anio':'2026','medida_unidad':'Personas','medida_valor':'100'}).status_code==302
        _,dashboard=context_get(company,'/dashboard?anio=2026')
        assert dashboard['total_emision']==pytest.approx(266.155,abs=0.02)
        assert dashboard['intensidad_emisiones']==pytest.approx(2.66155)
        _,history=context_get(company,'/historial?anio=2026&alcance=Alcance+1&q=Di%C3%A9sel')
        assert history['total']==1
        id=history['registros'][0]['id']
        export=company.get('/exportar_ghg?anio=2026')
        book=load_workbook(io.BytesIO(export.data),read_only=True,data_only=True)
        assert list(book['Resumen GHG'].values)[-1][1]==pytest.approx(266.155)
        book.close()
        other=client_for('empresa')
        with other.session_transaction() as session: session.update(user_id=4,empresa='AUDIT_B')
        assert other.get(f'/eliminar_cualquier_registro/manual/{id}').status_code==302
        _,unchanged=context_get(company,'/dashboard?anio=2026')
        assert unchanged['total_emision']==dashboard['total_emision']
        assert company.get(f'/eliminar_cualquier_registro/manual/{id}').status_code==302
        _,empty=context_get(company,'/dashboard?anio=2026')
        assert empty['total_emision']==0
        cursor=connection.cursor()
        for table in ['registros','combustible_movil','resultados_huella_suministro_agua']:
            cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE empresa='AUDIT_A'")
            assert cursor.fetchone()[0]==0


def water_post(client,tipo,volumen):
    return client.post('/registro',data={'alcance_oculto':'Agua','tipo_flujo':tipo,
        'periodo':'2026-01','volumen_m3':str(volumen),'retorna_mismo_sistema_hidrico':'si'})


def diagnosis(client,entrada='caudalimetro',salida='caudalimetro'):
    return client.post('/registro',data={'alcance_oculto':'Agua','solo_diagnostico_agua':'1',
        'medicion_entrada_agua':entrada,'medicion_salida_agua':salida})


def test_water_diagnosis_balance_dashboards_export_change_and_delete(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        company=client_for('empresa')
        assert diagnosis(company).status_code==204
        water_post(company,'captacion',100)
        water_post(company,'retorno',30)
        _,main=context_get(company,'/dashboard?anio=2026')
        _,water=context_get(company,'/agua?periodo=2026-01')
        assert main['huella_hidrica_total']==water['huella_hidrica_total']==70
        export=company.get('/agua/resultados/descargar?periodo=2026-01')
        book=load_workbook(io.BytesIO(export.data),read_only=True,data_only=True)
        values={row[0]:row[2] for row in list(book['Resumen'].values)[1:]}
        assert values['Huella hídrica total [m³]']==70
        book.close()
        diagnosis(company,'boleta','no')
        for path in ['/dashboard?anio=2026','/agua?periodo=2026-01']:
            _,context=context_get(company,path)
            assert context['huella_hidrica_total'] is None
            assert 'solo consumo de entrada' in context['mensaje_medicion_agua']
        diagnosis(company)
        assert company.post('/agua/registros/eliminar',data={'sede_id':'101','periodo':'2026-01','vista':'mensual'}).status_code==302
        _,water=context_get(company,'/agua?periodo=2026-01')
        assert water['consumo'] is None


def test_dashboards_show_supply_without_blue_footprint(acceptance_pg):
    connection,_=acceptance_pg
    cursor=connection.cursor()
    cursor.execute("""INSERT INTO registros (empresa,fecha,fuente,categoria,unidad,cantidad,factor,emision)
        VALUES ('AUDIT_A','2026-01-01','Electricidad','SEN','MWh',10,0,0)""")
    connection.commit()
    with database(connection):
        company=client_for('empresa')
        assert diagnosis(company,'no','no').status_code==204
        for path in ['/dashboard?anio=2026','/agua?periodo=2026-01']:
            response,context=context_get(company,path)
            html=response.get_data(as_text=True)
            assert 'No se puede obtener la huella azul' in html
            assert 'Total no disponible' in html
            assert context['huella_hidrica_total'] is None
            supply=context.get('huella_suministro_energetico',context.get('huella_suministro_total'))
            assert supply==pytest.approx(10.135227)
            assert context['suministro_disponible'] and context['suministro_completo']
        export=company.get('/agua/resultados/descargar?periodo=2026-01')
        book=load_workbook(io.BytesIO(export.data),read_only=True,data_only=True)
        values={row[0]:row[2] for row in list(book['Resumen'].values)[1:]}
        assert values['Huella de suministro [m³]']==pytest.approx(10.135227)
        assert values['Huella hídrica total [m³]']=='No disponible'
        book.close()


def test_admin_does_not_present_capture_only_as_measured_footprint(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        company=client_for('empresa')
        diagnosis(company,'boleta','no')
        water_post(company,'captacion',100)
        admin=client_for('admin')
        _,context=context_get(admin,'/admin/dashboard')
        index=context['chart_agua_empresas'].index('AUDIT_A')
        assert context['chart_agua_data']['Todos'][index] is None
        assert context['total_huella_hidrica_global'] is None
        _,coverage=context_get(admin,'/admin/huella-hidrica')
        assert all(row['resultado_calculable']!='Sí' for row in coverage['coberturas'] if row['empresa']=='AUDIT_A')


@pytest.mark.parametrize('value',['NaN','Infinity','-2'])
def test_invalid_productive_measure_rolls_back_configuration(acceptance_pg,value):
    connection,_=acceptance_pg
    with database(connection):
        client_for('empresa').post('/configuracion',data={'medida_anio':'2026','medida_unidad':'Personas','medida_valor':value})
    cursor=connection.cursor()
    cursor.execute("SELECT COUNT(*) FROM medidas_productivas WHERE empresa='AUDIT_A'")
    assert cursor.fetchone()[0]==0


def test_existing_integrity_constraints_upgrade_for_rename(pg):
    reconciliar_integridad(pg)
    cursor=pg.cursor()
    cursor.execute('ALTER TABLE combustible_movil DROP CONSTRAINT movil_consumo_empresa_fk')
    cursor.execute('ALTER TABLE combustible_movil ADD CONSTRAINT movil_consumo_empresa_fk FOREIGN KEY (registro_id,empresa) REFERENCES registros(id,empresa) ON DELETE CASCADE')
    pg.commit()
    reconciliar_integridad(pg)
    pg.commit()
    cursor.execute("UPDATE registros SET empresa='AUDIT_RENAMED' WHERE id=101")
    pg.commit()
    for table in ['combustible_movil','resultados_huella_suministro_agua']:
        cursor.execute(f'SELECT empresa FROM {table} WHERE registro_id=101')
        assert cursor.fetchone()[0]=='AUDIT_RENAMED'


@pytest.mark.parametrize('role',['admin','gestor_empresa'])
def test_create_edit_reset_and_login_company_without_granting_admin(acceptance_pg,role):
    connection,_=acceptance_pg
    password=secrets.token_urlsafe(18)
    replacement=secrets.token_urlsafe(18)
    with database(connection):
        manager=client_for(role)
        assert manager.post('/admin/crear_empresa',data={'empresa':'AUDIT_CREATED','email':'created@example.invalid',
            'password':password,'rol':'admin','es_admin':'1'}).status_code==302
        cursor=connection.cursor()
        cursor.execute("SELECT id,es_admin,rol FROM usuarios WHERE empresa='AUDIT_CREATED'")
        id,flag,created_role=cursor.fetchone()
        assert flag==0 and created_role in (None,'empresa')
        assert manager.post(f'/admin/editar_empresa/{id}',data={'empresa':'AUDIT_CREATED','email':'created@example.invalid','contacto':'Contacto prueba'}).status_code==302
        assert manager.post(f'/admin/resetear_password/{id}',data={'nueva_password':replacement}).status_code==302
        login=client_for()
        assert login.post('/api/login',json={'email':'created@example.invalid','password':password}).status_code==401
        assert login.post('/',data={'email':'created@example.invalid','password':replacement}).location=='/dashboard'


def test_rename_failure_rolls_back_all_related_tables(acceptance_pg):
    connection,_=acceptance_pg
    seed_consumption_and_support(connection)
    reconciliar_integridad(connection)
    connection.commit()
    with database(connection):
        response=client_for('admin').post('/admin/editar_empresa/3',data={'empresa':'AUDIT_RENAMED','email':'audit_b@example.invalid'})
        assert response.status_code==302
    cursor=connection.cursor()
    for table in ['usuarios','registros','combustible_movil','resultados_huella_suministro_agua','tickets','pending_pdf_uploads']:
        cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE empresa='AUDIT_A'")
        assert cursor.fetchone()[0]==1,table


def factor_payload(**overrides):
    return {'metodo':'M','version_metodo':'1','actividad':'Operación','nivel_geografico':'cuenca',
        'codigo_geografico':'C-01','periodo_inicio':'2026-01-01','periodo_fin':'2026-12-31',
        'factor_m3eq_m3':'0.5','fuente':'Prueba','referencia':'R',**overrides}


def test_admin_site_factor_and_toggle_update_company_measurement(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        admin=client_for('admin')
        assert admin.post('/admin/huella-hidrica/sedes',data={'sede_id':'101','empresa':'AUDIT_A',
            'nombre_sede':'Planta','region':'RM','comuna':'Santiago','codigo_cuenca':'C-01',
            'nombre_cuenca':'Cuenca','latitud':'-33','longitud':'-70','nivel_ubicacion':'cuenca'}).status_code==302
        assert admin.post('/admin/huella-hidrica/factores',data=factor_payload()).status_code==302
        company=client_for('empresa')
        diagnosis(company)
        water_post(company,'captacion',100)
        water_post(company,'retorno',30)
        _,water=context_get(company,'/agua?periodo=2026-01')
        assert water['huella']==35
        _,coverage=context_get(admin,'/admin/huella-hidrica')
        assert coverage['coberturas'][0]['resultado_calculable']=='Sí'
        cursor=connection.cursor()
        cursor.execute('SELECT id FROM factores_escasez_agua')
        id=cursor.fetchone()[0]
        assert admin.post(f'/admin/huella-hidrica/factores/{id}/toggle').status_code==302
        _,water=context_get(company,'/agua?periodo=2026-01')
        assert water['huella'] is None
        assert water['consumo']==70


@pytest.mark.parametrize('overrides',[{'factor_m3eq_m3':'NaN'},{'factor_m3eq_m3':'Infinity'},
    {'periodo_inicio':'2026-12-31','periodo_fin':'2026-01-01'}])
def test_admin_cannot_store_unusable_water_factor(acceptance_pg,overrides):
    connection,_=acceptance_pg
    with database(connection):
        assert client_for('admin').post('/admin/huella-hidrica/factores',data=factor_payload(**overrides)).status_code==302
    cursor=connection.cursor()
    cursor.execute('SELECT COUNT(*) FROM factores_escasez_agua')
    assert cursor.fetchone()[0]==0


def test_site_with_history_cannot_move_to_other_company(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        company=client_for('empresa')
        diagnosis(company)
        water_post(company,'captacion',100)
        assert client_for('admin').post('/admin/huella-hidrica/sedes',data={'sede_id':'101',
            'empresa':'AUDIT_B','nombre_sede':'Reasignada'}).status_code==302
    cursor=connection.cursor()
    cursor.execute('SELECT empresa FROM agua_sedes WHERE id=101')
    assert cursor.fetchone()[0]=='AUDIT_A'


def test_configuration_can_save_flags_without_productive_measure(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        response=client_for('empresa').post('/configuracion',data={'medida_anio':'2026',
            'medida_unidad':'','medida_valor':'','vehiculos':'si'})
        assert response.status_code==302
    cursor=connection.cursor()
    cursor.execute("SELECT vehiculos FROM configuracion WHERE empresa='AUDIT_A'")
    assert cursor.fetchone()[0]==1


@pytest.mark.parametrize('password',[None,123,''])
def test_api_login_bad_password_type_is_rejected(acceptance_pg,password):
    connection,_=acceptance_pg
    with database(connection):
        response=client_for().post('/api/login',json={'email':'audit_a@example.invalid','password':password})
        assert response.status_code in (400,401)


def test_admin_api_excel_and_dashboard_totals_agree(acceptance_pg):
    connection,password=acceptance_pg
    seed_consumption_and_support(connection)
    with database(connection):
        company=client_for('empresa')
        company.post('/configuracion',data={'medida_anio':'2026','medida_unidad':'Personas','medida_valor':'100'})
        admin=client_for('admin')
        token=admin.post('/api/login',json={'email':'audit_admin@example.invalid','password':password}).json['access_token']
        response=admin.get('/api/admin/emisiones-por-empresa?anio=2026',headers={'Authorization':'Bearer '+token})
        assert response.status_code==200
        _,dashboard=context_get(admin,'/admin/dashboard')
        assert sum(row['total_ghg_kgco2e'] for row in response.json)==dashboard['total_emisiones']==536
        row=next(r for r in response.json if r['empresa']=='AUDIT_A')
        assert row['medida_productiva_anio']['emisiones_anio_kgco2e']==268
        assert row['medida_productiva_anio']['intensidad_emisiones_kgco2e_unidad']==2.68
        export=admin.get('/admin/exportar_ghg_global?anio=Todos')
        assert export.status_code==200
        book=load_workbook(io.BytesIO(export.data),read_only=True,data_only=True)
        rows=list(book['Resumen Global'].values)
        column=rows[0].index('TOTAL GHG (kg CO₂e)')
        assert sum(row[column] for row in rows[1:])==536
        book.close()


def test_admin_does_not_average_water_intensities_of_different_units(acceptance_pg):
    connection,_=acceptance_pg
    with database(connection):
        for id,name,unit in [(3,'AUDIT_A','Personas'),(4,'AUDIT_B','Toneladas producto')]:
            company=client_for('empresa')
            with company.session_transaction() as session: session.update(user_id=id,empresa=name)
            diagnosis(company)
            water_post(company,'captacion',100)
            water_post(company,'retorno',30)
            company.post('/configuracion',data={'medida_anio':'2026','medida_unidad':unit,'medida_valor':'100'})
        for path in ['/admin/dashboard','/admin/empresas']:
            _,context=context_get(client_for('admin'),path)
            assert context['intensidad_huella_promedio'] is None
